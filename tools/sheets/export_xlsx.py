"""Export game data to an .xlsx workbook for viewing / importing into Google Sheets.

Reads the real data (items.json, sample loadouts, BattleConfig via the engine CLI) so the
sheet never drifts from the game. Computed columns are formulas, so edits in the sheet
recalculate. The sheet is a VIEW: the source of truth is still items.json.

Usage (repo root):
    cd server && ./mvnw -q compile && ./mvnw -q -pl battle-engine dependency:build-classpath -Dmdep.outputFile=target/cp.txt
    py tools/sheets/export_xlsx.py [--runs 500] [--out exports/One-Bag-Jianghu-data.xlsx]
"""
import argparse
import datetime
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
ENGINE = ROOT / "server" / "battle-engine"
ITEMS_JSON = ENGINE / "src" / "main" / "resources" / "data" / "items.json"
SAMPLES = ENGINE / "samples"
SECTS_JSON = ENGINE / "src" / "main" / "resources" / "data" / "sects.json"
UNITY_DATA = ROOT / "One-Bag-Jianghu" / "Assets" / "_Game" / "Resources" / "Data"
CHAPTERS_JSON = UNITY_DATA / "Pve" / "chapters.json"
LEVELS_JSON = UNITY_DATA / "levels.json"
POWER_JSON = UNITY_DATA / "power.json"
BATTLE_CONFIG_JAVA = ENGINE / "src" / "main" / "java" / "com" / "onebagjianghu" / "battle" / "sim" / "BattleConfig.java"

sys.path.insert(0, str(ROOT / "tools" / "balance"))
import power  # noqa: E402

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="3B2C35")
TITLE_FONT = Font(name=FONT, size=14, bold=True)
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT, color="0000FF")      # values copied from game data
FORMULA_FONT = Font(name=FONT, color="000000")   # computed
NOTE_FONT = Font(name=FONT, italic=True, color="666666")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# ---- design tables (source: docs/design/*.md) ----
RARITIES = [  # code, name, placeholder frame color; multipliers come from BattleConfig
    ("common", "Phàm Phẩm", "9E9E9E"),
    ("uncommon", "Lương Phẩm", "4CAF50"),
    ("rare", "Tinh Phẩm", "2F80ED"),
    ("epic", "Tuyệt Phẩm", "9B51E0"),
    ("legendary", "Thần Binh", "F2B01E"),
    ("supreme", "Chí Tôn", "E53935"),
]
SECTS = [  # code, name, weapon classes, playstyle (items.md)
    ("kim_cang", "Kim Cang Tự", "Quyền, Côn", "Cương Khí, Định Thân"),
    ("thai_thanh", "Thái Thanh Quán", "Kiếm, Phiến", "Nội lực, Hóa Kình"),
    ("liet_diem", "Liệt Diễm Đường", "Đao, Chưởng", "Liệt Hỏa"),
    ("huyen_doc", "Huyền Độc Môn", "Ám Khí, Trảo", "Kịch Độc, Xuất Huyết"),
    ("han_son", "Hàn Sơn Phái", "Tiêu, Thương", "Điểm Huyệt, Trì Trệ"),
    ("thiet_phu", "Thiết Phủ Bang", "Rìu, Chùy", "Phá Cương, sát thương lớn"),
    ("bich_thuy", "Bích Thủy Am", "Tiên, Cầm", "Liệu Thương, Điều Tức"),
]
WEAPON_CLASSES = [  # code, name, size, speed, mechanic (items.md)
    ("fist", "Quyền Sáo", "1x2", "Rất nhanh (1–1.5s)", "Liên Kích, ra đòn nhanh nhất"),
    ("palm", "Chưởng", "1x2", "Trung bình", "Nội Thương"),
    ("claw", "Trảo", "1x2", "Nhanh", "Xuất Huyết"),
    ("sword", "Kiếm", "1x3", "Trung bình", "Hội Tâm cao, Tuyệt Kỹ nhiều nhát"),
    ("blade", "Đao", "1x3", "Chậm", "Sát thương đơn lớn, Liệt Hỏa"),
    ("dual", "Song Đao / Song Kiếm", "2x2", "Nhanh", "Liên Kích x2"),
    ("spear", "Thương", "1x4", "Trung bình", "Xuyên Cương Khí"),
    ("staff", "Côn / Bổng", "1x4", "Trung bình", "Định Thân, cắt Tuyệt Kỹ"),
    ("axe", "Rìu (Phủ)", "2x2", "Rất chậm", "Phá Cương cực mạnh"),
    ("hammer", "Chùy", "2x2", "Chậm", "Định Thân, hạ Tốc Xuất Chiêu địch"),
    ("whip", "Tiên (roi)", "1x3", "Nhanh", "Trì Trệ binh khí địch"),
    ("fan", "Phiến", "1x2", "Trung bình", "Điểm Huyệt"),
    ("brush", "Phán Quan Bút", "1x2", "Trung bình", "Điểm Huyệt"),
    ("zither", "Cầm", "2x2", "Chậm", "Bỏ qua Cương Khí, hồi phục"),
    ("flute", "Tiêu", "1x2", "Chậm", "Bỏ qua Cương Khí, Trì Trệ toàn bộ"),
    ("hidden", "Ám Khí", "1x2", "Nhanh", "Tích Kịch Độc"),
    ("bow", "Cung / Nỏ", "1x3", "Chậm", "Nhiều phát, phá Đan Dược địch"),
]
STATUSES = [  # code, name, kind, mechanic (glossary.md / combat.md)
    ("burn", "Liệt Hỏa", "Tích điểm", "Đủ ngưỡng → bộc phát, Cương Khí đỡ được, rồi reset về 0"),
    ("poison", "Kịch Độc", "Tích điểm", "Đủ ngưỡng → bộc phát xuyên Cương Khí, rồi reset về 0"),
    ("bleed", "Xuất Huyết", "Tích điểm", "Đủ ngưỡng → bộc phát xuyên Cương Khí + % khí huyết tối đa, rồi reset"),
    ("regen", "Điều Tức", "Cộng dồn", "Mỗi giây hồi khí huyết bằng số tầng"),
    ("stun", "Định Thân", "Thời hạn", "Dừng mọi vật phẩm"),
    ("slow", "Trì Trệ", "Thời hạn", "Tốc xuất chiêu còn 50%"),
    ("seal", "Điểm Huyệt", "Thời hạn", "Dừng thời gian hồi của một món"),
    ("haste", "Khinh Công", "Thời hạn", "Tốc xuất chiêu x2"),
]
STATUS_NAME = {c: n for c, n, _, _ in STATUSES}
TYPE_TABS = [("weapon", "Binh Khí"), ("armor", "Hộ Cụ"), ("manual", "Công Pháp"), ("pill", "Đan Dược"), ("relic", "Kỳ Trân")]
TYPE_NAME = dict(TYPE_TABS)
CLASS_NAME = {c[0]: c[1] for c in WEAPON_CLASSES}
SECT_NAME = {s[0]: s[1] for s in SECTS}
RARITY_NAME = {r[0]: r[1] for r in RARITIES}

TARGET_NAME = {
    "self": "bản thân", "enemy": "đối thủ", "self_item": "món này", "event_source": "món vừa kích hoạt",
    "self_all_items": "mọi món phe mình", "self_other_items": "các món khác phe mình",
    "self_all_weapons": "mọi binh khí phe mình", "adjacent_items": "các món liền kề",
    "adjacent_weapons": "binh khí liền kề", "enemy_random_item": "1 món ngẫu nhiên của địch",
    "enemy_random_weapon": "1 binh khí ngẫu nhiên của địch", "enemy_all_items": "toàn bộ đồ địch",
}
EVENT_NAME = {
    "battle_start": "Vào trận", "hit": "Khi đánh trúng", "crit": "Khi Hội Tâm", "ultimate": "Khi tung Tuyệt Kỹ",
    "parry_success": "Khi Hóa Kình thành công", "parried": "Khi bị Hóa Kình", "hp_below": "Khi khí huyết dưới",
    "status_applied": "Khi gây trạng thái", "burst": "Khi trạng thái bộc phát",
    "damaged": "Khi chịu sát thương", "healed": "Khi được hồi khí huyết",
    "controlled": "Khi bị khống chế", "crit_taken": "Khi chịu Hội Tâm",
    "enemy_ultimate": "Khi đối thủ tung Tuyệt Kỹ", "qi_empty": "Khi cạn Nội Lực",
    "shield_broken": "Khi phá hết Cương Khí địch", "burst_taken": "Khi chịu Bộc Phát",
}
FROM_NAME = {"self": " (chính món này)", "adjacent": " (món liền kề)", "enemy": " (từ đối thủ)"}


def fmt(n):
    return f"{n:g}"


def describe_effect(e):
    op = e["op"].lower()
    tgt = e.get("target")
    t = f" lên {TARGET_NAME.get(tgt.lower(), tgt)}" if tgt else ""
    v = e.get("value", 0)
    pct = e.get("pct", 1.0)
    if op == "damage":
        amt = str(v) if v else f"{fmt(pct * 100)}% Lực Đạo"
        hits = e.get("hits", 1)
        s = f"Sát thương {hits} nhát × {amt}" if hits > 1 else f"Sát thương {amt}"
        if e.get("ignoreShield"):
            s += " (xuyên Cương Khí)"
        return s + (" vào bản thân" if tgt and tgt.lower() == "self" else "")
    if op == "gain_qi":
        return f"+{v} nội lực"
    if op == "drain_qi":
        return f"đối thủ −{v} nội lực"
    if op == "shield":
        return f"+{v} Cương Khí"
    if op == "heal":
        return f"hồi {v} khí huyết"
    if op == "shield_break":
        return f"Phá Cương {v}"
    if op == "dmg_reduce":
        return f"Hộ Thân +{fmt(pct * 100)}%"
    if op == "reduce_max_hp":
        return f"Nội Thương: đối thủ −{v} khí huyết tối đa"
    if op == "parry_discount":
        return f"Hóa Kình tốn ít hơn {v} nội lực"
    if op == "apply":
        st = e["status"].lower()
        name = STATUS_NAME.get(st, st)
        if st in ("burn", "poison", "bleed"):
            return f"+{e['stacks']} điểm {name}"
        if st == "regen":
            return f"+{e['stacks']} Điều Tức"
        return f"{name} {fmt(e['duration'])}s{t}"
    if op == "charge":
        return f"Vận Khí {fmt(e['duration'])}s{t or ' lên món này'}"
    if op == "add_attack":
        return f"+{v} Lực Đạo{t}"
    if op == "add_crit":
        return f"+{fmt(pct * 100)}% Hội Tâm{t}"
    if op == "buff_stack":
        stat = {"attack": "Công", "dmg_reduce": "Hộ Thân", "cooldown_reduction": "Giảm Hồi Chiêu",
                "crit": "Hội Tâm"}.get(str(e.get("stat", "")).lower(), e.get("stat", ""))
        amount = f"{fmt(pct * 100)}%" if e.get("pct") is not None else str(v)
        return f"Cộng dồn {stat} +{amount}{t} (tối đa {e.get('maxStacks', 1)} tầng)"
    if op == "cleanse":
        return f"Xóa Trì Trệ và Điểm Huyệt{t}"
    if op == "lifesteal":
        return f"Hút khí huyết {fmt(pct * 100)}%{t}"
    if op == "raise_max_hp":
        return f"Tăng khí huyết tối đa +{v}{t}"
    return op


def describe_effects(effects):
    return ", ".join(describe_effect(e) for e in effects)


def describe_trigger(t):
    on = t["on"].lower()
    head = EVENT_NAME.get(on, on)
    if on == "hp_below":
        head += f" {fmt(t['hpBelow'] * 100)}%"
    if t.get("status"):
        head += f" ({STATUS_NAME.get(t['status'].lower(), t['status'])})"
    head += FROM_NAME.get(str(t.get("from", "any")).lower(), "")
    s = f"{head}: {describe_effects(t.get('effects', []))}"
    if t.get("maxUses"):
        s += f" (tối đa {t['maxUses']} lần)"
    if t.get("minQiAfter") is not None:
        s += f" (sau khi trả còn ít nhất {t['minQiAfter']} Nội Lực)"
    return s


def describe_active(item):
    active = item.get("active")
    if not active:
        return ""
    effects = active.get("effects", active) if isinstance(active, dict) else active
    text = describe_effects(effects if isinstance(effects, list) else [])
    cooldown = item.get("cooldown")
    return f"{text} · Hồi chiêu {fmt(cooldown)} giây" if cooldown else text


def describe_unlock_stats(u):
    parts = []
    if u.get("attack"):
        parts.append(f"Lực Đạo {u['attack']:+d}")
    if u.get("cooldown"):
        parts.append(f"Tốc Xuất Chiêu {u['cooldown']:+g}s")
    if u.get("crit"):
        parts.append(f"Hội Tâm {u['crit'] * 100:+g}%")
    if u.get("ultimateCost"):
        parts.append(f"Chi phí Tuyệt Kỹ {u['ultimateCost']:+d}")
    return ", ".join(parts)


def describe_tam_phap(lv):
    parts = []
    if lv.get("attackPct"):
        parts.append(f"Lực Đạo +{lv['attackPct'] * 100:g}%")
    if lv.get("cooldownPct"):
        parts.append(f"Tốc Xuất Chiêu −{lv['cooldownPct'] * 100:g}%")
    if lv.get("crit"):
        parts.append(f"Hội Tâm +{lv['crit'] * 100:g}%")
    if lv.get("ultimateCost"):
        parts.append(f"Chi phí Tuyệt Kỹ {lv['ultimateCost']:+d}")
    if lv.get("basic"):
        parts.append("Phổ Công thêm: " + describe_effects(lv["basic"]))
    if lv.get("ultimate"):
        parts.append("Tuyệt Kỹ thêm: " + describe_effects(lv["ultimate"]))
    parts += [describe_trigger(t) for t in lv.get("triggers", [])]
    return "; ".join(parts)


def describe_unlock(u):
    parts = []
    stats = describe_unlock_stats(u)
    if stats:
        parts.append(stats)
    if u.get("basic"):
        parts.append("Phổ Công thêm: " + describe_effects(u["basic"]))
    if u.get("ultimate"):
        parts.append("Tuyệt Kỹ thêm: " + describe_effects(u["ultimate"]))
    parts += [describe_trigger(t) for t in u.get("triggers", [])]
    return "; ".join(parts)


def java(args):
    cp_file = ENGINE / "target" / "cp.txt"
    if not cp_file.exists():
        raise SystemExit("Thiếu target/cp.txt — chạy lệnh build ở đầu file trước.")
    cp = f"{ENGINE / 'target' / 'classes'};{cp_file.read_text().strip()}"
    return subprocess.run(["java", "-cp", cp, "com.onebagjianghu.battle.cli.BattleCli", *args],
                          capture_output=True, text=True, encoding="utf-8", check=True).stdout


def load_config(override_path=None):
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "config.json"
        java(["--dump-config", str(out)])
        config = json.loads(out.read_text(encoding="utf-8"))
    if "rarityMultiplierPermille" not in config:
        source = BATTLE_CONFIG_JAVA.read_text(encoding="utf-8")
        match = re.search(r"rarityMultiplierPermille\s*=\s*\{([^}]+)\}", source)
        if not match:
            raise ValueError("Không tìm thấy rarityMultiplierPermille trong config hoặc BattleConfig.java")
        config["rarityMultiplierPermille"] = [int(v) for v in re.findall(r"\d+", match.group(1))]
    if override_path:
        override = json.loads(Path(override_path).read_text(encoding="utf-8"))
        def merge(base, extra):
            for key, value in extra.items():
                base[key] = merge(base.get(key, {}), value) if isinstance(value, dict) and isinstance(base.get(key), dict) else value
            return base
        config = merge(config, override)
    return config


def engine_options(items_path, sects_path, battle_config):
    args = ["--items", str(items_path), "--sects", str(sects_path)]
    if battle_config:
        args += ["--battle-config", str(battle_config)]
    return args


def load_matrix(runs, loadouts_path, items_path, sects_path, battle_config):
    text = java(["--matrix", str(loadouts_path), "--runs", str(runs), *engine_options(items_path, sects_path, battle_config)])
    rows = []
    for line in text.splitlines():
        m = re.match(r"#(\d+) (.+?)\s{2,}(.*)$", line)
        if m:
            cells = m.group(3).split()
            values = [None if c == "-" else float(c.rstrip("%")) / 100 for c in cells]
            rows.append((m.group(2).strip(), values[:-1], values[-1]))
    return rows


def load_fighter_stats(loadouts, items_path, sects_path, battle_config):
    result = {}
    paths = [p for p, _ in loadouts]
    if not paths:
        return result
    pattern = re.compile(r"^\[A\].*?maxHp (\d+) · Hộ Thân (\d+)‰ · Giảm Hồi Chiêu (\d+)‰", re.M)
    for path in paths:
        text = java([str(path), str(paths[0]), "--runs", "2", *engine_options(items_path, sects_path, battle_config)])
        match = pattern.search(text)
        if not match:
            raise ValueError(f"Không đọc được fighterStats từ Breakdown cho {path}")
        result[path.stem] = tuple(map(int, match.groups()))
    return result


# ---- sheet helpers ----

def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT


def header(ws, row, names, widths):
    for i, (name, width) in enumerate(zip(names, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def put(ws, row, col, value, formula=False, fmt_code=None, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = FORMULA_FONT if formula else DATA_FONT
    c.border = BORDER
    if fmt_code:
        c.number_format = fmt_code
    if wrap:
        c.alignment = WRAP
    return c


def rarity_fill(cell, code):
    color = {r[0]: r[2] for r in RARITIES}.get(code)
    if color:
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")


def display_value(value):
    """Keep scalar JSON values readable and preserve nested/unknown fields."""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(", ", ": "))
    return value


def extra_fields(obj, known):
    return "; ".join(f"{key}: {display_value(value)}" for key, value in obj.items() if key not in known)


def item_stat(item):
    kind = item["type"].lower()
    stats = item.get("stats", {})
    if kind == "weapon":
        return item.get("attack", ""), None
    key = {"armor": "hp", "relic": "hp", "manual": "cooldownReduction", "pill": "dmgReduce"}.get(kind)
    value = stats.get(key, "") if key else ""
    return value, "0%" if key in ("cooldownReduction", "dmgReduce") and value != "" else None


# ---- build ----

def build(runs, out_path, items_path=ITEMS_JSON, sects_path=SECTS_JSON, battle_config=None,
          loadouts_path=SAMPLES, chapters_path=CHAPTERS_JSON, levels_path=LEVELS_JSON):
    items_path, sects_path, loadouts_path = map(Path, (items_path, sects_path, loadouts_path))
    chapters_path, levels_path = Path(chapters_path), Path(levels_path)
    items = json.loads(items_path.read_text(encoding="utf-8"))
    sects = json.loads(sects_path.read_text(encoding="utf-8"))
    buffed_by = lambda it: ", ".join(s["name"] for s in sects if it.get("weaponClass") in s["weaponClasses"])
    config = load_config(battle_config)
    rarity_mults = config["rarityMultiplierPermille"]
    if len(rarity_mults) != len(RARITIES):
        raise ValueError(f"Cần {len(RARITIES)} hệ số phẩm, nhận được {len(rarity_mults)}")
    materials = json.loads((UNITY_DATA / "materials.json").read_text(encoding="utf-8"))
    upgrade_costs = json.loads((UNITY_DATA / "upgrade_costs.json").read_text(encoding="utf-8"))
    merge = json.loads((UNITY_DATA / "merge.json").read_text(encoding="utf-8"))
    shop = json.loads((UNITY_DATA / "shop.json").read_text(encoding="utf-8"))
    rewards = json.loads((UNITY_DATA / "rewards.json").read_text(encoding="utf-8"))
    starter = json.loads((UNITY_DATA / "starter.json").read_text(encoding="utf-8"))
    loadouts = [(p, json.loads(p.read_text(encoding="utf-8"))) for p in sorted(loadouts_path.glob("*.json"))]
    matrix = load_matrix(runs, loadouts_path, items_path, sects_path, battle_config)
    fighter_stats = load_fighter_stats(loadouts, items_path, sects_path, battle_config)
    chapters = json.loads(chapters_path.read_text(encoding="utf-8"))
    levels = json.loads(levels_path.read_text(encoding="utf-8"))
    power_data = json.loads(POWER_JSON.read_text(encoding="utf-8"))
    today = datetime.date.today().isoformat()

    wb = Workbook()

    # Hằng Số first (other tabs reference it), then moved to its display position at the end.
    hs = wb.active
    hs.title = "Hằng Số"
    title(hs, "Hằng số chiến đấu", f"Nguồn: BattleConfig.java (xuất {today}). Đổi ở đây KHÔNG đổi game — sửa trong code/dữ liệu.")
    header(hs, 4, ["Hằng số", "Giá trị", "Đơn vị", "Ghi chú"], [34, 12, 12, 60])
    consts = [
        ("Tick", config["tickMs"], "ms", "Bước thời gian mô phỏng"),
        ("Nội lực tối đa", config["maxQi"], "", ""),
        ("Chi phí Hóa Kình", config["parryCost"], "nội lực", "Giảm được bằng parry_discount"),
        ("Chi phí Hóa Kình tối thiểu", config["minParryCost"], "nội lực", ""),
        ("Định Thân khi bị Hóa Kình", config["parryStunMs"] / 1000, "giây", ""),
        ("Kiệt Lực bắt đầu", config["exhaustionStartMs"] / 1000, "giây", "Giây thứ n mất n × bước máu"),
        ("Bước Kiệt Lực", config["exhaustionStep"], "máu", ""),
        ("Điều Tức tối đa", config["maxRegen"], "mỗi giây", ""),
        ("Hết giờ (hòa)", config["maxDurationMs"] / 1000, "giây", ""),
        ("Trần Cương Khí", config["shieldCapPermille"] / 1000, "% Sinh Lực tối đa", "Cương Khí vượt trần bị bỏ"),
        ("Hộ Thân tối đa", config["maxDmgReducePermille"] / 1000, "", ""),
        ("Giảm Hồi Chiêu tối đa", config["maxCooldownReducePermille"] / 1000, "", ""),
        ("Khí Huyết cơ bản", config["baseHp"], "máu", ""),
        ("Khí Huyết mỗi cấp", config["hpPerLevel"], "máu", ""),
        ("Ngân sách cơ bản", 6, "ST/giây/ô", "Quy tắc 1 trong docs/design/balance.md"),
        ("Hao hụt khi tích điểm", 0.25, "", "Điểm chưa đủ ngưỡng khi hết trận (balance.md)"),
        ("Hệ số xuyên Cương Khí", 1.2, "", "Quy đổi giá trị của sát thương xuyên khiên (balance.md)"),
    ]
    ref = {}
    r = 5
    for name, value, unit, note in consts:
        put(hs, r, 1, name).font = Font(name=FONT)
        pct = name in ("Trần Cương Khí", "Hộ Thân tối đa", "Giảm Hồi Chiêu tối đa", "Hao hụt khi tích điểm")
        put(hs, r, 2, value, fmt_code="0%" if pct else None)
        put(hs, r, 3, unit).font = Font(name=FONT)
        put(hs, r, 4, note).font = NOTE_FONT
        ref[name] = f"'Hằng Số'!$B${r}"
        r += 1

    r += 1
    hs.cell(row=r, column=1, value="Độ hiếm").font = Font(name=FONT, bold=True)
    r += 1
    for i, h in enumerate(["Mã", "Tên", "Hệ số chỉ số", "Màu khung (tạm)"], start=1):
        c = hs.cell(row=r, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    rar_first = r + 1
    for (code, name, color), mult in zip(RARITIES, rarity_mults):
        r += 1
        put(hs, r, 1, code)
        c = put(hs, r, 2, name)
        rarity_fill(c, code)
        put(hs, r, 3, mult / 1000, fmt_code="0.00")
        put(hs, r, 4, "#" + color)
    rar_last = r
    ref["rar_code"] = f"'Hằng Số'!$A${rar_first}:$A${rar_last}"
    ref["rar_name"] = f"'Hằng Số'!$B${rar_first}:$B${rar_last}"
    ref["rar_mult"] = f"'Hằng Số'!$C${rar_first}:$C${rar_last}"
    hs.cell(row=r + 1, column=1, value="Màu khung là màu tạm; sẽ đổi theo phong cách art đã chọn (docs/art/style-options.md).").font = NOTE_FONT

    r += 3
    hs.cell(row=r, column=1, value="Trạng thái tích điểm (bộc phát)").font = Font(name=FONT, bold=True)
    r += 1
    for i, h in enumerate(["Mã", "Tên", "Ngưỡng", "ST bộc phát", "% KH tối đa", "Xuyên Cương Khí", "Tự giảm/giây", "ST quy đổi mỗi điểm"], start=1):
        c = hs.cell(row=r, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        hs.column_dimensions[get_column_letter(i)].width = max(hs.column_dimensions[get_column_letter(i)].width or 0, 14)
    burst_first = r + 1
    for code in ("burn", "poison", "bleed"):
        b = config[code]
        r += 1
        put(hs, r, 1, code)
        put(hs, r, 2, STATUS_NAME[code])
        put(hs, r, 3, b["threshold"])
        put(hs, r, 4, b["damage"])
        put(hs, r, 5, b["maxHpPermille"] / 1000, fmt_code="0%")
        put(hs, r, 6, "Có" if b["ignoreShield"] else "Không")
        put(hs, r, 7, b["decayPerSecond"])
        put(hs, r, 8, f'=(D{r}+E{r}*{ref["Khí Huyết cơ bản"]})/C{r}*IF(F{r}="Có",{ref["Hệ số xuyên Cương Khí"]},1)',
            formula=True, fmt_code="0.00")
    burst_last = r
    ref["burst_code"] = f"'Hằng Số'!$A${burst_first}:$A${burst_last}"
    ref["burst_value"] = f"'Hằng Số'!$H${burst_first}:$H${burst_last}"

    r += 2
    hs.cell(row=r, column=1, value="Hệ số Trọng").font = Font(name=FONT, bold=True)
    r += 1
    for i, h in enumerate(["Trọng", "Hệ số"], start=1):
        c = hs.cell(row=r, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    tier_first = r + 1
    for tier, mult in enumerate(config["tierMultiplierPermille"], start=1):
        r += 1
        put(hs, r, 1, tier)
        put(hs, r, 2, mult / 1000, fmt_code="0.00")
    tier_last = r
    ref["tier_mult"] = f"'Hằng Số'!$B${tier_first}:$B${tier_last}"

    r += 2
    hs.cell(row=r, column=1, value="Khí Huyết cơ bản theo cấp").font = Font(name=FONT, bold=True)
    r += 1
    for col, name in enumerate(("Cấp", "Khí Huyết cơ bản"), 1):
        c = hs.cell(row=r, column=col, value=name); c.font = HEADER_FONT; c.fill = HEADER_FILL
    for level in range(1, 31):
        r += 1
        put(hs, r, 1, level)
        put(hs, r, 2, f"={ref['Khí Huyết cơ bản']}+{ref['Khí Huyết mỗi cấp']}*(A{r}-1)", formula=True)

    r += 2
    hs.cell(row=r, column=1, value="Chỉ số cuối theo Trọng × Phẩm").font = Font(name=FONT, bold=True)
    hs.cell(row=r + 1, column=1, value="Chỉ số gốc minh họa").font = Font(name=FONT)
    put(hs, r + 1, 2, next(it["attack"] for it in items if it.get("attack", 0) > 0))
    r += 2
    for col, (_, name, _) in enumerate(RARITIES, start=2):
        c = hs.cell(row=r, column=col, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    c = hs.cell(row=r, column=1, value="Trọng")
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    final_header = r
    for tier in range(1, len(config["tierMultiplierPermille"]) + 1):
        r += 1
        put(hs, r, 1, tier)
        for col in range(2, 2 + len(RARITIES)):
            put(hs, r, col, f"=INT(INT($B${final_header - 1}*INDEX({ref['tier_mult']},$A{r}))*INDEX({ref['rar_mult']},COLUMN()-1))", formula=True)
    hs.freeze_panes = None

    # ---- Tất Cả Vật Phẩm (lookup table for other tabs) ----
    all_ws = wb.create_sheet("Tất Cả Vật Phẩm")
    title(all_ws, "Tất cả vật phẩm", "Bảng tra cứu dùng chung. Dữ liệu (chữ xanh) lấy từ items.json; chữ đen là công thức.")
    header(all_ws, 4, ["ID", "Tên", "Loại", "Tâm Pháp buff", "Độ hiếm", "Rộng", "Cao", "Số ô", "Chỉ số", "Chỉ số T9 Chí Tôn"], [24, 24, 12, 18, 14, 8, 8, 8, 14, 20])
    all_first = 5
    for i, it in enumerate(items):
        row = all_first + i
        put(all_ws, row, 1, it["id"])
        put(all_ws, row, 2, it["name"])
        put(all_ws, row, 3, TYPE_NAME[it["type"].lower()])
        put(all_ws, row, 4, buffed_by(it))
        rarity = it.get("rarity", "")
        rarity_fill(put(all_ws, row, 5, RARITY_NAME.get(rarity.lower(), rarity) if rarity else ""), rarity.lower())
        put(all_ws, row, 6, it.get("size", [1, 1])[0])
        put(all_ws, row, 7, it.get("size", [1, 1])[1])
        put(all_ws, row, 8, f"=F{row}*G{row}", formula=True)
        stat, stat_fmt = item_stat(it)
        put(all_ws, row, 9, stat, fmt_code=stat_fmt)
        scale_base = f"I{row}*1000" if stat_fmt else f"I{row}"
        scale_tail = "/1000" if stat_fmt else ""
        put(all_ws, row, 10, f'=IF(I{row}="","",INT(INT({scale_base}*INDEX({ref["tier_mult"]},9))*INDEX({ref["rar_mult"]},6)){scale_tail})', formula=True, fmt_code=stat_fmt)
    all_last = all_first + len(items) - 1
    lk = lambda col: f"'Tất Cả Vật Phẩm'!${col}${all_first}:${col}${all_last}"

    # ---- Binh Khí ----
    ws = wb.create_sheet("Binh Khí")
    title(ws, "Binh Khí", "Chữ xanh: số liệu từ items.json. Chữ đen: công thức. ST/giây chưa tính Tuyệt Kỹ và hiệu ứng khống chế; xem docs/design/balance.md.")
    cols = ["ID", "Tên", "Hệ", "Tâm Pháp buff", "Độ hiếm", "Rộng", "Cao", "Số ô", "Chỉ số", "Tốc Xuất Chiêu (giây)",
            "Hội Tâm", "Nội lực mỗi đòn", "Trạng thái tích", "Điểm mỗi đòn", "Tuyệt Kỹ (nội lực)",
            "Phổ Công", "Tuyệt Kỹ", "ST/giây từ đòn", "ST/giây từ trạng thái", "Tổng ST/giây", "Ngân sách ST/giây", "% ngân sách",
            "Đột Phá Trọng 3", "Đột Phá Trọng 6", "Đột Phá Trọng 9", "Chỉ số T9 Chí Tôn"]
    header(ws, 4, cols, [22, 20, 12, 16, 12, 7, 7, 7, 9, 11, 9, 10, 12, 9, 11, 40, 40, 11, 12, 11, 11, 11, 34, 34, 34, 18])
    weapons = [it for it in items if it["type"].lower() == "weapon"]
    for i, it in enumerate(weapons):
        row = 5 + i
        basic = it.get("basic", [])
        qi = sum(e.get("value", 0) for e in basic if e["op"].lower() == "gain_qi")
        st = next((e for e in basic if e["op"].lower() == "apply" and e["status"].lower() in ("burn", "poison", "bleed")), None)
        ult = it.get("ultimate")
        put(ws, row, 1, it["id"])
        put(ws, row, 2, it["name"])
        put(ws, row, 3, CLASS_NAME.get(it.get("weaponClass"), it.get("weaponClass")))
        put(ws, row, 4, buffed_by(it))
        rarity_fill(put(ws, row, 5, RARITY_NAME[it["rarity"].lower()]), it["rarity"].lower())
        put(ws, row, 6, it["size"][0])
        put(ws, row, 7, it["size"][1])
        put(ws, row, 8, f"=F{row}*G{row}", formula=True)
        put(ws, row, 9, it["attack"])
        put(ws, row, 10, it["cooldown"], fmt_code="0.0")
        put(ws, row, 11, it.get("crit", 0), fmt_code="0%")
        put(ws, row, 12, qi)
        put(ws, row, 13, st["status"].lower() if st else "")
        put(ws, row, 14, st["stacks"] if st else 0)
        put(ws, row, 15, ult["cost"] if ult else "")
        put(ws, row, 16, describe_effects(basic), wrap=True)
        put(ws, row, 17, describe_effects(ult["effects"]) if ult else "", wrap=True)
        put(ws, row, 18, f"=I{row}/J{row}*(1+K{row})", formula=True, fmt_code="0.0")
        put(ws, row, 19, f'=IF(M{row}="",0,N{row}/J{row}*INDEX({ref["burst_value"]},MATCH(M{row},{ref["burst_code"]},0))*(1-{ref["Hao hụt khi tích điểm"]}))',
            formula=True, fmt_code="0.0")
        put(ws, row, 20, f"=R{row}+S{row}", formula=True, fmt_code="0.0")
        put(ws, row, 21, f'={ref["Ngân sách cơ bản"]}*H{row}*INDEX({ref["rar_mult"]},MATCH(E{row},{ref["rar_name"]},0))',
            formula=True, fmt_code="0.0")
        put(ws, row, 22, f"=IF(U{row}=0,0,T{row}/U{row})", formula=True, fmt_code="0%")
        by_tier = {u["tier"]: u for u in it.get("unlocks", [])}
        for k, tier in enumerate((3, 6, 9)):
            put(ws, row, 23 + k, describe_unlock(by_tier[tier]) if tier in by_tier else "", wrap=True)
        put(ws, row, 26, f"=INT(INT(I{row}*INDEX({ref['tier_mult']},9))*INDEX({ref['rar_mult']},6))", formula=True)
    last = 4 + len(weapons)
    ws.conditional_formatting.add(f"V5:V{last}", ColorScaleRule(
        start_type="num", start_value=0.6, start_color="F8696B",
        mid_type="num", mid_value=1.0, mid_color="FFFFFF",
        end_type="num", end_value=1.4, end_color="63BE7B"))
    ws.cell(row=last + 2, column=1, value="% ngân sách: ~100% là đúng chuẩn. Món có Tuyệt Kỹ mạnh / khống chế thường thấp hơn 100% vì phần đó chưa được tính.").font = NOTE_FONT

    # ---- Đột Phá ----
    ws = wb.create_sheet("Đột Phá")
    title(ws, "Đột Phá", "Khi vật phẩm lên Trọng 3 / 6 / 9 thì mở thêm hiệu ứng, cộng dồn các mốc trước. Số liệu từ trường unlocks trong items.json. Giá trị hiệu ứng (sát thương, khiên, điểm trạng thái…) vẫn nhân hệ số Trọng; nội lực và thời gian thì không.")
    header(ws, 4, ["ID", "Tên", "Loại", "Trọng", "Thay đổi chỉ số", "Phổ Công thêm", "Tuyệt Kỹ thêm", "Kích hoạt thêm"],
           [22, 22, 12, 8, 34, 34, 34, 50])
    row = 5
    for it in items:
        for u in it.get("unlocks", []):
            put(ws, row, 1, it["id"])
            put(ws, row, 2, it["name"])
            put(ws, row, 3, TYPE_NAME.get(it["type"].lower(), it["type"]))
            put(ws, row, 4, u["tier"])
            put(ws, row, 5, describe_unlock_stats(u), wrap=True)
            put(ws, row, 6, describe_effects(u.get("basic", [])), wrap=True)
            put(ws, row, 7, describe_effects(u.get("ultimate", [])), wrap=True)
            put(ws, row, 8, "\n".join(describe_trigger(t) for t in u.get("triggers", [])), wrap=True)
            row += 1
    if row == 5:
        ws.cell(row=5, column=1, value="(chưa có Đột Phá)").font = NOTE_FONT

    # ---- passive item tabs ----
    for type_code, tab in TYPE_TABS[1:]:
        ws = wb.create_sheet(tab)
        title(ws, tab, "Chữ xanh: số liệu từ items.json. Hiệu ứng được mô tả tự động từ dữ liệu.")
        header(ws, 4, ["ID", "Tên", "Tâm Pháp buff", "Độ hiếm", "Rộng", "Cao", "Số ô", "Tốc Xuất Chiêu (giây)", "Chỉ số", "Chỉ số T9 Chí Tôn", "Chiêu chủ động", "Hiệu ứng"],
               [22, 22, 16, 12, 7, 7, 7, 12, 14, 20, 50, 80])
        rows = [it for it in items if it["type"].lower() == type_code]
        for i, it in enumerate(rows):
            row = 5 + i
            put(ws, row, 1, it["id"])
            put(ws, row, 2, it["name"])
            put(ws, row, 3, buffed_by(it))
            rarity_fill(put(ws, row, 4, RARITY_NAME[it["rarity"].lower()]), it["rarity"].lower())
            put(ws, row, 5, it.get("size", [1, 1])[0])
            put(ws, row, 6, it.get("size", [1, 1])[1])
            put(ws, row, 7, f"=E{row}*F{row}", formula=True)
            put(ws, row, 8, it.get("cooldown", 0) or "", fmt_code="0.0")
            parts = []
            if it.get("basic"):
                parts.append("Mỗi lần kích hoạt: " + describe_effects(it["basic"]))
            parts += [describe_trigger(t) for t in it.get("triggers", [])]
            stat, stat_fmt = item_stat(it)
            put(ws, row, 9, stat, fmt_code=stat_fmt)
            scale_base = f"I{row}*1000" if stat_fmt else f"I{row}"
            scale_tail = "/1000" if stat_fmt else ""
            put(ws, row, 10, f'=IF(I{row}="","",INT(INT({scale_base}*INDEX({ref["tier_mult"]},9))*INDEX({ref["rar_mult"]},6)){scale_tail})', formula=True, fmt_code=stat_fmt)
            put(ws, row, 11, describe_active(it), wrap=True)
            put(ws, row, 12, "\n".join(parts), wrap=True)
        if not rows:
            ws.cell(row=5, column=1, value="(chưa có vật phẩm)").font = NOTE_FONT

    # ---- Hiệp Khách + Đội Hình Mẫu ----
    hk = wb.create_sheet("Hiệp Khách", 0)
    dh = wb.create_sheet("Đội Hình Mẫu", 1)
    title(dh, "Đội hình mẫu (từng món)", "Nguồn: server/battle-engine/samples/*.json. Tên và số ô tra từ tab Tất Cả Vật Phẩm.")
    header(dh, 4, ["Đội hình", "ID vật phẩm", "Tên", "Loại", "x", "y", "Trọng", "Số ô"], [16, 24, 24, 12, 6, 6, 8, 8])
    row = 5
    for path, lo in loadouts:
        key = path.stem
        for p in lo["items"]:
            put(dh, row, 1, key)
            put(dh, row, 2, p["id"])
            put(dh, row, 3, f'=INDEX({lk("B")},MATCH(B{row},{lk("A")},0))', formula=True)
            put(dh, row, 4, f'=INDEX({lk("C")},MATCH(B{row},{lk("A")},0))', formula=True)
            put(dh, row, 5, p["x"])
            put(dh, row, 6, p["y"])
            put(dh, row, 7, p.get("tier", 1))
            put(dh, row, 8, f'=INDEX({lk("H")},MATCH(B{row},{lk("A")},0))', formula=True)
            row += 1
    dh_last = row - 1

    title(hk, "Hiệp Khách", "Cấp lấy từ đội hình; Khí Huyết tối đa, Hộ Thân và Giảm Hồi Chiêu lấy trực tiếp từ BattleSimulator.fighterStats qua CLI Breakdown.")
    header(hk, 4, ["Đội hình", "Hiệp Khách", "Cấp", "Khí Huyết tối đa", "Hộ Thân", "Giảm Hồi Chiêu", "Bao Phục (cột)", "Bao Phục (hàng)", "Tổng ô",
                   "Ô đã dùng", "% lấp đầy", "Số binh khí", "Tỉ lệ thắng TB", "Môn phái", "Cấp Tâm Pháp"],
           [16, 36, 8, 14, 11, 16, 10, 10, 8, 10, 10, 10, 12, 18, 10])
    avg_col = get_column_letter(3 + len(loadouts))
    for i, (path, lo) in enumerate(loadouts):
        key = path.stem
        row = 5 + i
        max_hp, dmg_reduce, cdr = fighter_stats[key]
        put(hk, row, 1, key)
        put(hk, row, 2, lo["name"])
        put(hk, row, 3, lo.get("level", "") if "level" in lo else "")
        put(hk, row, 4, max_hp)
        put(hk, row, 5, dmg_reduce / 1000, fmt_code="0.0%")
        put(hk, row, 6, cdr / 1000, fmt_code="0.0%")
        put(hk, row, 7, lo.get("cols", 5)); put(hk, row, 8, lo.get("rows", 4))
        put(hk, row, 9, f"=G{row}*H{row}", formula=True)
        put(hk, row, 10, f"=SUMIF('Đội Hình Mẫu'!$A$5:$A${dh_last},A{row},'Đội Hình Mẫu'!$H$5:$H${dh_last})", formula=True)
        put(hk, row, 11, f"=J{row}/I{row}", formula=True, fmt_code="0%")
        put(hk, row, 12, f"=COUNTIFS('Đội Hình Mẫu'!$A$5:$A${dh_last},A{row},'Đội Hình Mẫu'!$D$5:$D${dh_last},\"Binh Khí\")", formula=True)
        put(hk, row, 13, f"=IFERROR(INDEX('Đối Đầu'!${avg_col}$5:${avg_col}${4 + len(loadouts)},MATCH(B{row},'Đối Đầu'!$A$5:$A${4 + len(loadouts)},0)),\"\")", formula=True, fmt_code="0.0%")
        put(hk, row, 14, SECT_NAME.get(lo.get("sect") or "", "")); put(hk, row, 15, lo.get("tamPhapLevel", 1))

    # ---- Trạng Thái, Môn Phái, Hệ Binh Khí ----
    ws = wb.create_sheet("Trạng Thái")
    title(ws, "Trạng thái", "Nguồn: docs/design/glossary.md, combat.md. Thông số bộc phát ở tab Hằng Số.")
    header(ws, 4, ["Mã", "Tên", "Loại", "Cơ chế"], [10, 14, 12, 80])
    for i, s in enumerate(STATUSES):
        for j, v in enumerate(s, start=1):
            put(ws, 5 + i, j, v, wrap=(j == 4))

    ws = wb.create_sheet("Môn Phái")
    title(ws, "Môn phái và Tâm Pháp", "Nguồn: server/battle-engine/src/main/resources/data/sects.json. Vật phẩm trung lập; môn phái chỉ tác động qua Tâm Pháp, buff các hệ binh khí liệt kê. Các mốc cộng dồn.")
    header(ws, 4, ["Mã", "Môn phái", "Tâm Pháp", "Hệ binh khí được buff", "Lối chơi", "Cấp 1", "Cấp 3", "Cấp 6", "Cấp 9",
                   "Số binh khí được buff"], [12, 18, 24, 22, 26, 30, 30, 30, 30, 12])
    style_of = {code: style for code, _, _, style in SECTS}
    class_name = {c[0]: c[1] for c in WEAPON_CLASSES}
    for i, sd in enumerate(sects):
        row = 5 + i
        put(ws, row, 1, sd["id"])
        put(ws, row, 2, sd["name"])
        put(ws, row, 3, sd["tamPhap"])
        put(ws, row, 4, ", ".join(class_name.get(c, c) for c in sd["weaponClasses"]))
        put(ws, row, 5, style_of.get(sd["id"], ""))
        by_level = {lv["level"]: lv for lv in sd.get("levels", [])}
        for k, level in enumerate((1, 3, 6, 9)):
            put(ws, row, 6 + k, describe_tam_phap(by_level[level]) if level in by_level else "", wrap=True)
        put(ws, row, 10, f'=COUNTIF(\'Binh Khí\'!$D$5:$D$200,"*"&B{row}&"*")', formula=True)

    ws = wb.create_sheet("Hệ Binh Khí")
    title(ws, "Hệ binh khí", "Nguồn: docs/design/items.md. Binh khí tối thiểu 2 ô.")
    header(ws, 4, ["Mã", "Hệ", "Kích thước", "Tốc độ", "Cơ chế đặc trưng", "Số vật phẩm hiện có"], [10, 20, 10, 18, 36, 12])
    for i, (code, name, size, speed, mech) in enumerate(WEAPON_CLASSES):
        row = 5 + i
        for j, v in enumerate((code, name, size, speed, mech), start=1):
            put(ws, row, j, v)
        put(ws, row, 6, f"=COUNTIF('Binh Khí'!$C$5:$C$200,B{row})", formula=True)

    # ---- Economy data (Unity Resources/Data) ----
    material_names = {m["id"]: m["name"] for m in materials}
    item_names = {it["id"]: it["name"] for it in items}

    ws = wb.create_sheet("Nguyên Liệu")
    title(ws, "Nguyên liệu", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/materials.json")
    material_keys = ["id", "name", "itemType", "band", "tier", "isCatalyst", "pvpOnly", "description"]
    header(ws, 4, ["ID", "Tên", "Loại vật phẩm", "Nhóm", "Bậc", "Chất dẫn?", "Chỉ PvP?", "Mô tả", "Trường khác"],
           [25, 25, 16, 12, 8, 11, 11, 45, 35])
    for i, material in enumerate(materials, start=5):
        for col, key in enumerate(material_keys, start=1):
            put(ws, i, col, display_value(material.get(key, "")), wrap=(key == "description"))
        put(ws, i, 9, extra_fields(material, material_keys), wrap=True)

    ws = wb.create_sheet("Nâng Trọng")
    title(ws, "Nâng Trọng", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/upgrade_costs.json. Tổng Bạc là công thức cộng các bước T1→T9 theo phẩm.")
    step_keys = ["fromTier", "materialTier", "materialCount", "catalystCount", "silver"]
    header(ws, 4, ["Mã phẩm", "Phẩm", "Từ Trọng", "Lên Trọng", "Bậc nguyên liệu", "Số nguyên liệu", "Số chất dẫn", "Bạc", "Tổng Bạc T1→T9", "Trường khác"],
           [14, 15, 10, 10, 14, 14, 12, 12, 18, 35])
    row = 5
    upgrade_first = row
    for rarity_code, steps in upgrade_costs.get("steps", {}).items():
        for step_index, step in enumerate(steps):
            put(ws, row, 1, rarity_code)
            rarity_fill(put(ws, row, 2, RARITY_NAME.get(rarity_code, rarity_code)), rarity_code)
            put(ws, row, 3, step.get("fromTier", ""))
            put(ws, row, 4, step.get("fromTier", 0) + 1 if "fromTier" in step else "")
            put(ws, row, 5, step.get("materialTier", ""))
            put(ws, row, 6, step.get("materialCount", ""))
            put(ws, row, 7, step.get("catalystCount", ""))
            put(ws, row, 8, step.get("silver", ""), fmt_code="#,##0")
            if step_index == 0:
                put(ws, row, 9, f'=SUMIF($A$5:$A${4 + sum(len(v) for v in upgrade_costs.get("steps", {}).values())},A{row},$H$5:$H${4 + sum(len(v) for v in upgrade_costs.get("steps", {}).values())})', formula=True, fmt_code="#,##0")
            put(ws, row, 10, extra_fields(step, step_keys), wrap=True)
            row += 1
    upgrade_last = row - 1
    if extra_fields(upgrade_costs, {"steps"}):
        put(ws, upgrade_last + 2, 1, "Trường cấp cao khác")
        put(ws, upgrade_last + 2, 2, extra_fields(upgrade_costs, {"steps"}), wrap=True)

    ws = wb.create_sheet("Hợp Thành")
    title(ws, "Hợp Thành", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/merge.json")
    merge_keys = ["fromRarity", "toRarity", "silver", "extraMaterials"]
    header(ws, 4, ["Mã phẩm từ", "Phẩm từ", "Mã phẩm đến", "Phẩm đến", "Bạc", "Nguyên liệu thêm", "Trường khác"],
           [14, 15, 14, 15, 12, 42, 35])
    for row, step in enumerate(merge.get("steps", []), start=5):
        from_code, to_code = step.get("fromRarity", ""), step.get("toRarity", "")
        put(ws, row, 1, from_code)
        rarity_fill(put(ws, row, 2, RARITY_NAME.get(from_code, from_code)), from_code)
        put(ws, row, 3, to_code)
        rarity_fill(put(ws, row, 4, RARITY_NAME.get(to_code, to_code)), to_code)
        put(ws, row, 5, step.get("silver", ""), fmt_code="#,##0")
        extras = step.get("extraMaterials", {})
        put(ws, row, 6, "; ".join(f"{material_names.get(mid, mid)} ×{count}" for mid, count in extras.items()), wrap=True)
        put(ws, row, 7, extra_fields(step, merge_keys), wrap=True)
    if extra_fields(merge, {"steps"}):
        put(ws, 5 + len(merge.get("steps", [])) + 1, 1, "Trường cấp cao khác")
        put(ws, 5 + len(merge.get("steps", [])) + 1, 2, extra_fields(merge, {"steps"}), wrap=True)

    ws = wb.create_sheet("Thương Hội")
    title(ws, "Thương Hội", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/shop.json")
    put(ws, 4, 1, "Thiết lập làm mới").font = Font(name=FONT, bold=True)
    header(ws, 5, ["Giờ làm mới miễn phí", "Chi phí làm mới trả phí", "Trường khác"], [22, 38, 45])
    put(ws, 6, 1, shop.get("freeRefreshHours", ""))
    put(ws, 6, 2, display_value(shop.get("paidRefreshCosts", [])))
    known_shop = {"freeRefreshHours", "paidRefreshCosts", "itemBasePrices", "slots", "exchange", "reputationOffers"}
    put(ws, 6, 3, extra_fields(shop, known_shop), wrap=True)
    put(ws, 8, 1, "Giá vật phẩm cơ bản theo phẩm").font = Font(name=FONT, bold=True)
    header(ws, 9, ["Mã phẩm", "Phẩm", "Giá cơ bản"], [14, 15, 15])
    row = 10
    for code, price in shop.get("itemBasePrices", {}).items():
        put(ws, row, 1, code)
        rarity_fill(put(ws, row, 2, RARITY_NAME.get(code, code)), code)
        put(ws, row, 3, price, fmt_code="#,##0")
        row += 1
    row += 1
    put(ws, row, 1, "Ô hàng (lưới Thương Hội)").font = Font(name=FONT, bold=True)
    row += 1
    slot_header = row
    slot_keys = ["type", "rarities", "weights", "tier", "quantity", "price", "discountPercent"]
    header(ws, row, ["Ô", "Loại", "Phẩm (trọng số)", "Trọng/Bậc", "Số lượng", "Giá", "Giảm giá", "Trường khác"],
           [8, 20, 42, 12, 12, 12, 12, 35])
    for slot_no, slot in enumerate(shop.get("slots", []), start=1):
        row += 1
        put(ws, row, 1, slot_no)
        put(ws, row, 2, slot.get("type", ""))
        pairs = [f"{RARITY_NAME.get(code, code)}: {weight}" for code, weight in zip(slot.get("rarities", []), slot.get("weights", []))]
        put(ws, row, 3, "; ".join(pairs), wrap=True)
        put(ws, row, 4, slot.get("tier", ""))
        put(ws, row, 5, slot.get("quantity", ""))
        put(ws, row, 6, slot.get("price", ""), fmt_code="#,##0")
        put(ws, row, 7, slot.get("discountPercent", "") / 100 if "discountPercent" in slot else "", fmt_code="0%")
        put(ws, row, 8, extra_fields(slot, slot_keys), wrap=True)
    row += 2
    put(ws, row, 1, "Đổi nguyên liệu").font = Font(name=FONT, bold=True)
    row += 1
    exchange = shop.get("exchange", {})
    exchange_keys = list(exchange.keys())
    header(ws, row, exchange_keys or ["(không có)"], [22] * max(1, len(exchange_keys)))
    row += 1
    for col, key in enumerate(exchange_keys, start=1):
        put(ws, row, col, display_value(exchange[key]))
    row += 2
    put(ws, row, 1, "Đổi Danh Vọng").font = Font(name=FONT, bold=True)
    row += 1
    offers = shop.get("reputationOffers", [])
    offer_keys = list(dict.fromkeys(key for offer in offers for key in offer))
    offer_labels = {"id": "ID", "name": "Tên", "price": "Giá Danh Vọng", "limit": "Giới hạn", "quantity": "Số lượng"}
    header(ws, row, [offer_labels.get(key, key) for key in offer_keys] or ["(không có)"], [25] * max(1, len(offer_keys)))
    for offer in offers:
        row += 1
        for col, key in enumerate(offer_keys, start=1):
            put(ws, row, col, display_value(offer.get(key, "")), wrap=True)

    ws = wb.create_sheet("Phần Thưởng")
    title(ws, "Phần thưởng", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/rewards.json")
    reward_rows = [(key, value) for key, value in rewards.items() if isinstance(value, dict)]
    reward_keys = list(dict.fromkeys(key for _, reward in reward_rows for key in reward))
    reward_names = {"win": "Thắng", "loss": "Thua", "firstClear": "Vượt lần đầu"}
    header(ws, 4, ["Kết quả"] + reward_keys, [18] + [20] * len(reward_keys))
    for row, (code, reward) in enumerate(reward_rows, start=5):
        put(ws, row, 1, reward_names.get(code, code))
        for col, key in enumerate(reward_keys, start=2):
            value = reward.get(key, "")
            fmt_code = "0.0%" if key.endswith("Permille") and value != "" else None
            put(ws, row, col, value / 1000 if fmt_code else display_value(value), fmt_code=fmt_code)
    scalar_rewards = {key: value for key, value in rewards.items() if not isinstance(value, dict)}
    if scalar_rewards:
        put(ws, 6 + len(reward_rows), 1, "Trường cấp cao khác")
        put(ws, 6 + len(reward_rows), 2, display_value(scalar_rewards), wrap=True)

    ws = wb.create_sheet("Khởi Đầu")
    title(ws, "Khởi đầu", "Nguồn: One-Bag-Jianghu/Assets/_Game/Resources/Data/starter.json")
    put(ws, 4, 1, "Tiền tệ ban đầu").font = Font(name=FONT, bold=True)
    currency_keys = list(starter.get("currencies", {}).keys())
    header(ws, 5, currency_keys or ["(không có)"], [20] * max(1, len(currency_keys)))
    for col, key in enumerate(currency_keys, start=1):
        put(ws, 6, col, starter["currencies"][key])
    put(ws, 8, 1, "Vật phẩm và vị trí trong Bao Phục").font = Font(name=FONT, bold=True)
    starter_items = starter.get("items", [])
    starter_loadout = starter.get("loadout", [])
    item_extra_keys = list(dict.fromkeys(k for obj in starter_items for k in obj if k not in {"key", "itemId", "tier", "rarity"}))
    pos_extra_keys = list(dict.fromkeys(k for obj in starter_loadout for k in obj if k not in {"itemKey", "x", "y"}))
    header(ws, 9, ["Khóa", "ID vật phẩm", "Tên", "Trọng", "Phẩm", "x", "y"] + item_extra_keys + pos_extra_keys,
           [16, 26, 26, 10, 15, 8, 8] + [20] * (len(item_extra_keys) + len(pos_extra_keys)))
    positions = {p.get("itemKey"): p for p in starter_loadout}
    for row, starter_item in enumerate(starter_items, start=10):
        pos = positions.get(starter_item.get("key"), {})
        code = starter_item.get("rarity", "")
        values = [starter_item.get("key", ""), starter_item.get("itemId", ""), item_names.get(starter_item.get("itemId"), starter_item.get("itemId", "")),
                  starter_item.get("tier", ""), RARITY_NAME.get(code, code), pos.get("x", ""), pos.get("y", "")]
        for col, value in enumerate(values, start=1):
            cell = put(ws, row, col, value)
            if col == 5:
                rarity_fill(cell, code)
        for col, key in enumerate(item_extra_keys, start=8):
            put(ws, row, col, display_value(starter_item.get(key, "")), wrap=True)
        for col, key in enumerate(pos_extra_keys, start=8 + len(item_extra_keys)):
            put(ws, row, col, display_value(pos.get(key, "")), wrap=True)
    unmatched_positions = [p for p in starter_loadout if p.get("itemKey") not in {it.get("key") for it in starter_items}]
    for pos in unmatched_positions:
        row = ws.max_row + 1
        put(ws, row, 1, pos.get("itemKey", ""))
        put(ws, row, 6, pos.get("x", ""))
        put(ws, row, 7, pos.get("y", ""))
    other_starter = extra_fields(starter, {"currencies", "items", "loadout"})
    if other_starter:
        put(ws, ws.max_row + 2, 1, "Trường cấp cao khác")
        put(ws, ws.max_row, 2, other_starter, wrap=True)

    # ---- PvE ----
    item_names = {it["id"]: it["name"] for it in items}
    ws = wb.create_sheet("Hành Tẩu")
    title(ws, "Hành Tẩu Giang Hồ",
          f"Nguồn: {chapters_path}. Lực Chiến quái do công cụ tính bằng power.lc().")
    put(ws, 3, 1, "Tỷ lệ thưởng đánh lại (%)").font = Font(name=FONT, bold=True)
    put(ws, 3, 2, chapters.get("repeatRewardPercent", 0))
    header(ws, 4, ["Hồi", "Ải", "Tên", "Loại", "Thể Lực", "Địch", "Cấp", "Khí Huyết (file cũ)",
                   "Môn phái", "Vật phẩm địch", "Phần thưởng", "Lực Chiến quái", "Lực Chiến đề nghị",
                   "Số món quái", "Đánh lại: Bạc", "Đánh lại: EXP"],
           [8, 8, 26, 12, 10, 28, 8, 18, 18, 70, 70, 18, 20, 14, 18, 18])
    row = 5
    for chapter in chapters.get("chapters", []):
        for stage in chapter.get("stages", []):
            enemy = stage.get("enemy", {})
            enemy_items = []
            for entry in enemy.get("items", []) or []:
                code = entry.get("id", "")
                rarity = entry.get("rarity", "")
                enemy_items.append(f"{item_names.get(code, code)} · T{entry.get('tier', 1)} · {RARITY_NAME.get(rarity, rarity)}")
            values = [chapter.get("index", ""), stage.get("index", ""), stage.get("name", ""), stage.get("kind", ""),
                      stage.get("staminaCost", ""), enemy.get("name", ""), enemy.get("level", ""), enemy.get("hp", ""),
                      SECT_NAME.get(enemy.get("sect") or "", enemy.get("sect") or ""), "\n".join(enemy_items),
                      display_value(stage.get("rewards", {}))]
            for col, value in enumerate(values, 1): put(ws, row, col, value, wrap=col in (10, 11))
            enemy_lc = power.lc(enemy, power_data)
            put(ws, row, 12, enemy_lc)
            put(ws, row, 13, f"=ROUND(L{row}*'Lực Chiến'!$B$7,0)", formula=True)
            put(ws, row, 14, len(enemy.get("items", []) or []))
            first_clear = stage.get("rewards", {}).get("firstClear", {})
            silver = first_clear.get("silver", 0)
            hero_exp = first_clear.get("heroExp", 0)
            silver_formula = f"=MAX(1,ROUND({silver}*$B$3/100,0))" if silver > 0 else "=0"
            exp_formula = f"=MAX(1,ROUND({hero_exp}*$B$3/100,0))" if hero_exp > 0 else "=0"
            put(ws, row, 15, silver_formula, formula=True)
            put(ws, row, 16, exp_formula, formula=True)
            row += 1
    row += 2
    put(ws, row, 1, "Chủ đề từng Hồi").font = Font(name=FONT, bold=True); row += 1
    theme_keys = list(dict.fromkeys(k for ch in chapters.get("chapters", []) for k in ch.get("theme", {})))
    header(ws, row, ["Hồi"] + theme_keys, [10] + [22] * len(theme_keys))
    for chapter in chapters.get("chapters", []):
        row += 1; put(ws, row, 1, chapter.get("index", ""))
        for col, key in enumerate(theme_keys, 2): put(ws, row, col, chapter.get("theme", {}).get(key, ""))
    row += 2; put(ws, row, 1, "Vật phẩm quái").font = Font(name=FONT, bold=True)
    header(ws, row + 1, ["ID", "Tên"], [28, 30])
    for code in chapters.get("monsterItems", []):
        row += 1; put(ws, row + 1, 1, code); put(ws, row + 1, 2, item_names.get(code, code))

    ws = wb.create_sheet("Cấp Độ")
    title(ws, "Cấp độ Hiệp Khách", f"Nguồn: {levels_path}. Khí Huyết cơ bản là công thức tham chiếu tab Hằng Số.")
    header(ws, 4, ["Cấp", "EXP tới cấp sau", "EXP tích lũy", "Mở khóa", "Khí Huyết cơ bản"], [10, 18, 18, 28, 20])
    max_level = levels.get("maxLevel", len(levels.get("expToNext", [])) + 1)
    unlocks = levels.get("unlocks", {})
    for level in range(1, max_level + 1):
        row = 4 + level
        put(ws, row, 1, level)
        exp = levels.get("expToNext", [])
        put(ws, row, 2, exp[level - 1] if level <= len(exp) else "")
        put(ws, row, 3, "=0" if level == 1 else f"=C{row-1}+B{row-1}", formula=True)
        put(ws, row, 4, unlocks.get(str(level), ""))
        put(ws, row, 5, f"={ref['Khí Huyết cơ bản']}+{ref['Khí Huyết mỗi cấp']}*(A{row}-1)", formula=True)

    # ---- Lực Chiến ----
    ws = wb.create_sheet("Lực Chiến")
    title(ws, "Lực Chiến",
          "Các số được hiệu chỉnh bằng mô phỏng (quyết định #51). Game đọc power.json; sheet này chỉ là bản xem.")
    put(ws, 4, 1, "Công thức").font = Font(name=FONT, bold=True)
    put(ws, 4, 2,
        "Lực Chiến = tổng (điểm gốc món × hệ số Trọng × hệ số phẩm) "
        "+ levelBase + levelPoints × (cấp − 1)", wrap=True)
    for row, (label, key) in enumerate((("levelBase", "levelBase"), ("levelPoints", "levelPoints"),
                                        ("recommendedFactor", "recommendedFactor")), start=5):
        put(ws, row, 1, label)
        put(ws, row, 2, power_data[key])
    header(ws, 9, ["Trọng", "tierFactor", "", "Phẩm", "rarityFactor"], [12, 16, 3, 20, 16])
    for row, (tier, factor) in enumerate(power_data["tierFactor"].items(), start=10):
        put(ws, row, 1, f"T{tier}")
        put(ws, row, 2, factor)
    for row, (rarity, factor) in enumerate(power_data["rarityFactor"].items(), start=10):
        put(ws, row, 4, RARITY_NAME.get(rarity, rarity))
        put(ws, row, 5, factor)
    item_by_id = {item["id"]: item for item in items}
    type_order = {code: index for index, (code, _) in enumerate(TYPE_TABS)}
    power_items = []
    for item_id, base in power_data["items"].items():
        item = item_by_id.get(item_id, {})
        item_type = item.get("type", "")
        power_items.append((type_order.get(item_type, len(type_order)), -base, item_id,
                            item.get("name", item_id), TYPE_NAME.get(item_type, item_type), base))
    power_items.sort()
    header(ws, 21, ["ID", "Tên", "Loại", "Điểm gốc", "LC T1 Phàm", "LC T9 Chí Tôn"],
           [30, 30, 18, 14, 16, 18])
    for row, (_, _, item_id, name, type_name, base) in enumerate(power_items, start=22):
        put(ws, row, 1, item_id)
        put(ws, row, 2, name)
        put(ws, row, 3, type_name)
        put(ws, row, 4, base)
        put(ws, row, 5, f"=ROUND(D{row}*$B$10*$E$10,0)", formula=True)
        put(ws, row, 6, f"=ROUND(D{row}*$B$18*$E$15,0)", formula=True)

    # ---- Đối Đầu ----
    ws = wb.create_sheet("Đối Đầu")
    title(ws, "Bảng đối đầu (tỉ lệ thắng của hàng trước cột)",
          f"Chạy {runs} trận mỗi cặp ngày {today} bằng battle-engine. Đây là KẾT QUẢ MÔ PHỎNG, không tự tính lại — chạy lại script để cập nhật. Hòa tính nửa trận thắng.")
    names = [n for n, _, _ in matrix]
    header(ws, 4, ["Đội hình"] + [f"#{i + 1}" for i in range(len(names))] + ["", "TB"], [36] + [10] * len(names) + [2, 10])
    for i, (name, values, avg) in enumerate(matrix):
        row = 5 + i
        put(ws, row, 1, f"{name}")
        for j, v in enumerate(values):
            put(ws, row, 2 + j, v if v is not None else "-", fmt_code="0.0%")
        put(ws, row, 3 + len(names), avg, fmt_code="0.0%")
    ws.conditional_formatting.add(f"B5:{get_column_letter(1 + len(names))}{4 + len(names)}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B",
        mid_type="num", mid_value=0.5, mid_color="FFFFFF",
        end_type="num", end_value=1, end_color="63BE7B"))
    ws.freeze_panes = "B5"

    # ---- Hướng Dẫn (first tab) ----
    ws = wb.create_sheet("Hướng Dẫn", 0)
    source_note = f"Xuất tự động ngày {today} bằng tools/sheets/export_xlsx.py"
    if items_path.resolve() != ITEMS_JSON.resolve():
        source_note += f". File dùng: items={items_path}; sects={sects_path}; loadouts={loadouts_path}; chapters={chapters_path}; levels={levels_path}"
    title(ws, "One Bag Jianghu — Dữ liệu game", source_note)
    guide = [
        ("Tab", "Nội dung"),
        ("Hiệp Khách", "Cấp, Khí Huyết tối đa, Hộ Thân, Giảm Hồi Chiêu từ engine; số ô và tỉ lệ thắng"),
        ("Đội Hình Mẫu", "Từng món trong từng đội hình mẫu, vị trí x/y, Trọng"),
        ("Binh Khí", "Chỉ số binh khí + ST/giây ước tính so với ngân sách theo ô (cân bằng), tóm tắt Đột Phá"),
        ("Đột Phá", "Hiệu ứng mở thêm khi vật phẩm lên Trọng 3 / 6 / 9"),
        ("Hộ Cụ / Công Pháp / Đan Dược / Kỳ Trân", "Vật phẩm bị động/hỗ trợ, hiệu ứng mô tả tự động"),
        ("Tất Cả Vật Phẩm", "Bảng tra cứu chung cho các tab khác"),
        ("Nguyên Liệu", "Danh mục nguyên liệu nâng Trọng, Chất dẫn và nguyên liệu PvP từ materials.json"),
        ("Nâng Trọng", "Chi phí từng bước nâng Trọng và tổng Bạc T1→T9 theo phẩm"),
        ("Hợp Thành", "Chi phí Hợp Thành lên phẩm kế tiếp và nguyên liệu bổ sung"),
        ("Thương Hội", "Làm mới, giá theo phẩm, ô hàng, đổi nguyên liệu và đổi Danh Vọng"),
        ("Phần Thưởng", "Phần thưởng thắng, thua và vượt lần đầu"),
        ("Khởi Đầu", "Tiền tệ, vật phẩm và vị trí Bao Phục ban đầu"),
        ("Hành Tẩu", "Mỗi ải PvE, đội hình địch, phần thưởng; chủ đề từng Hồi và vật phẩm quái"),
        ("Cấp Độ", "EXP, EXP tích lũy, mở khóa và Khí Huyết cơ bản theo cấp"),
        ("Lực Chiến", "Công thức, hệ số và điểm gốc vật phẩm từ power.json; các cột LC tự tính lại"),
        ("Môn Phái", "Môn phái và Tâm Pháp (nội tại buff các hệ binh khí), từ sects.json"),
        ("Trạng Thái / Hệ Binh Khí", "Bảng thiết kế"),
        ("Hằng Số", "Hằng số chiến đấu, phẩm, bộc phát trạng thái, hệ số Trọng và bảng Trọng × Phẩm"),
        ("Đối Đầu", "Tỉ lệ thắng giữa các đội hình mẫu (kết quả mô phỏng)"),
        ("", ""),
        ("Màu chữ", ""),
        ("Xanh dương", "Số liệu lấy từ dữ liệu game (items.json, samples, BattleConfig)"),
        ("Đen", "Công thức — tự tính lại khi sửa ô xanh"),
        ("", ""),
        ("Lưu ý", "Sheet này là BẢN XEM. Sửa số ở đây để thử nghiệm được, nhưng game chỉ đổi khi sửa items.json (sẽ có công cụ nhập ngược từ sheet sau)."),
    ]
    for i, (a, b) in enumerate(guide):
        ws.cell(row=4 + i, column=1, value=a).font = Font(name=FONT, bold=(i == 0 or a in ("Màu chữ", "Lưu ý")))
        ws.cell(row=4 + i, column=2, value=b).font = Font(name=FONT, bold=(i == 0))
        ws.cell(row=4 + i, column=2).alignment = WRAP
    ws.cell(row=4 + [a for a, _ in guide].index("Xanh dương"), column=1).font = Font(name=FONT, color="0000FF")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 100

    # final tab order
    order = ["Hướng Dẫn", "Hiệp Khách", "Đội Hình Mẫu", "Binh Khí", "Đột Phá", "Hộ Cụ", "Công Pháp", "Đan Dược", "Kỳ Trân",
             "Tất Cả Vật Phẩm", "Nguyên Liệu", "Nâng Trọng", "Hợp Thành", "Thương Hội", "Phần Thưởng", "Khởi Đầu",
             "Hành Tẩu", "Cấp Độ", "Lực Chiến", "Trạng Thái", "Môn Phái", "Hệ Binh Khí", "Hằng Số", "Đối Đầu"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(items), len(loadouts), len(matrix)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument("--runs", type=int, default=500)
    ap.add_argument("--out", default=str(ROOT / "exports" / "One-Bag-Jianghu-data.xlsx"))
    ap.add_argument("--items", default=str(ITEMS_JSON))
    ap.add_argument("--sects", default=str(SECTS_JSON))
    ap.add_argument("--battle-config")
    ap.add_argument("--loadouts", default=str(SAMPLES))
    ap.add_argument("--chapters", default=str(CHAPTERS_JSON))
    ap.add_argument("--levels", default=str(LEVELS_JSON))
    a = ap.parse_args()
    n_items, n_lo, n_mx = build(a.runs, Path(a.out), a.items, a.sects, a.battle_config, a.loadouts, a.chapters, a.levels)
    print(f"Đã xuất {n_items} vật phẩm, {n_lo} đội hình, bảng đối đầu {n_mx}x{n_mx} → {a.out}")
